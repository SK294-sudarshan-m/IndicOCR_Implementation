"""XLSX: one unit per sheet (cached formula values), then one OCR unit per embedded image."""

from __future__ import annotations

import datetime as dt
import zipfile
from pathlib import Path

from ..errors import CorruptFile
from ..schema import Unit
from .base import Context, md_table, ocr_embedded_image


def _fmt(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        return str(int(value)) if value == value and abs(value) < 1e15 and value == int(value) else repr(value)
    if isinstance(value, dt.datetime):
        return value.date().isoformat() if value.time() == dt.time(0) else value.isoformat(sep=" ")
    if isinstance(value, (dt.date, dt.time)):
        return value.isoformat()
    return str(value)


def _load(path: Path, data_only: bool):
    import openpyxl

    try:
        return openpyxl.load_workbook(path, data_only=data_only)
    except (zipfile.BadZipFile, KeyError, ValueError, OSError) as exc:
        raise CorruptFile(f"not a readable XLSX workbook: {exc}") from None
    except Exception as exc:  # openpyxl raises InvalidFileException and assorted XML errors
        raise CorruptFile(f"not a readable XLSX workbook: {type(exc).__name__}: {exc}") from None


def _cells(ws) -> dict[tuple[int, int], object]:
    """Non-empty cells only. ``ws._cells`` avoids walking the million blank-but-formatted rows a sheet can declare."""
    return {(r, c): cell.value for (r, c), cell in ws._cells.items() if cell.value is not None}


def extract(path: Path, ctx: Context) -> list[Unit]:
    from openpyxl.utils import get_column_letter

    values = _load(path, data_only=True)
    formulas = _load(path, data_only=False)
    units: list[Unit] = []
    for ws in values.worksheets:
        cells = _cells(ws)
        rows_n = max((r for r, _ in cells), default=0)
        cols_n = max((c for _, c in cells), default=0)
        grid = [[_fmt(cells.get((r, c))) for c in range(1, cols_n + 1)] for r in range(1, rows_n + 1)]

        warnings = []
        if ws.sheet_state != "visible":
            warnings.append(f"sheet is {ws.sheet_state}")
        formula_cells = _cells(formulas[ws.title])
        uncached = [
            f"{get_column_letter(c)}{r}"
            for (r, c), v in sorted(formula_cells.items())
            if isinstance(v, str) and v.startswith("=") and (r, c) not in cells
        ]
        if uncached:
            shown = ", ".join(uncached[:5]) + (" ..." if len(uncached) > 5 else "")
            warnings.append(f"{len(uncached)} formula cell(s) have no cached value (workbook was never calculated in Excel); shown empty: {shown}")

        body = md_table(grid) if grid else "_(empty sheet)_"
        units.append(
            Unit(
                0,
                "sheet",
                "native",
                f"worksheet, {rows_n} rows x {cols_n} columns, cached values read directly",
                text="\n".join("\t".join(row) for row in grid),
                markdown=f"## {ws.title}\n\n{body}",
                warnings=warnings,
                name=ws.title,
            )
        )
        for n, image in enumerate(getattr(ws, "_images", []), 1):
            anchor = getattr(getattr(image, "anchor", None), "_from", None)
            where = f" at row {anchor.row + 1}, column {anchor.col + 1}" if anchor is not None else ""
            try:
                blob = image._data()
            except Exception as exc:
                ctx.warnings.append(f"image {n} on sheet {ws.title!r} could not be read: {exc}")
                continue
            unit = ocr_embedded_image(ctx, blob, name=f"{ws.title}: image {n}", reason=f"image embedded in sheet {ws.title!r}{where}")
            if unit is not None:
                units.append(unit)
    return units
