"""Generates every test fixture from code (nothing is downloaded, nothing is checked in).

    python tests/make_fixtures.py <output-dir>

Writes the files plus truth.json (the ground-truth text for each fixture). Hindi is drawn with Nirmala UI
from C:\\Windows\\Fonts, shaped by MuPDF's HTML engine, so the pixels are proper Devanagari.
"""

from __future__ import annotations

import io
import json
import re
import sys
import zipfile
from pathlib import Path

import pymupdf

FONTS = Path("C:/Windows/Fonts")
NIRMALA = FONTS / "NIRMALA.TTF"
SEGOE_ICONS = FONTS / "segmdl2.ttf"  # has glyphs in the Private Use Area

EN_TITLE = "Quarterly Report"
EN_P1 = (
    "Reading text that is already machine-readable is faster and more accurate than running optical "
    "character recognition on a picture of it."
)
EN_P2 = "A document pipeline should decide, page by page, which route to take, and it should always record why."
HI_TITLE = "भारत की भाषाएँ"
HI_P1 = "भारत विविधताओं का देश है। यहाँ हर कुछ किलोमीटर पर बोली बदल जाती है और हर राज्य की अपनी संस्कृति है।"
HI_P2 = "हमारे संविधान में बाईस भाषाओं को मान्यता दी गई है।"
HI_SENTENCE = "भारत एक विशाल देश है। यहाँ अनेक भाषाएँ बोली जाती हैं।"
HI_PAGE_TEXT = " ".join([HI_TITLE, HI_P1, HI_P2])

_CSS = (
    "@font-face{font-family:nir;src:url(NIRMALA.TTF);} "
    "@font-face{font-family:nir;font-weight:bold;src:url(NIRMALAB.TTF);} body{font-family:nir;font-size:%dpt;}"
)


def fonts_available() -> bool:
    return NIRMALA.is_file() and SEGOE_ICONS.is_file()


def _shaped(html: str, size_pt: int = 13, width: float = 595, height: float = 842, box=(60, 60, 535, 500)):
    doc = pymupdf.open()
    page = doc.new_page(width=width, height=height)
    page.insert_htmlbox(pymupdf.Rect(*box), html, css=_CSS % size_pt, archive=pymupdf.Archive(str(FONTS)))
    return doc


def hindi_shaped_doc():
    return _shaped(f"<h2>{HI_TITLE}</h2><p>{HI_P1}</p><p>{HI_P2}</p>")


def english_doc():
    doc = pymupdf.open()
    for title, paragraph in ((EN_TITLE, EN_P1), (None, EN_P2)):
        page = doc.new_page(width=595, height=842)
        y = 80
        if title:
            page.insert_text((60, y), title, fontsize=22, fontname="hebo")
            y += 40
        page.insert_textbox(pymupdf.Rect(60, y, 535, y + 200), paragraph, fontsize=13, fontname="helv")
    return doc


def _png(page, dpi: int) -> bytes:
    return page.get_pixmap(dpi=dpi, alpha=False).tobytes("png")


def _degraded_jpeg(page, dpi: int = 300) -> bytes:
    """A poor scan: blur, noise, a slight skew, harsh JPEG compression."""
    import numpy as np
    from PIL import Image, ImageFilter

    pix = page.get_pixmap(dpi=dpi, alpha=False)
    im = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
    im = im.rotate(1.2, resample=Image.BICUBIC, fillcolor=(255, 255, 255)).filter(ImageFilter.GaussianBlur(1.2))
    noise = np.random.default_rng(7).normal(0, 12, (im.height, im.width, 1))
    im = Image.fromarray(np.clip(np.asarray(im).astype(np.float32) + noise, 0, 255).astype(np.uint8))
    buf = io.BytesIO()
    im.save(buf, "JPEG", quality=45)
    return buf.getvalue()


def image_only(src, path: Path, dpi: int = 200, hidden_text: str | None = None, hidden_font=None) -> None:
    """Rebuild ``src`` as pictures of its pages (no text layer). ``hidden_text`` adds an invisible text layer."""
    out = pymupdf.open()
    for n, page in enumerate(src):
        new = out.new_page(width=page.rect.width, height=page.rect.height)
        new.insert_image(new.rect, stream=_png(page, dpi))
        if hidden_text is not None and n == 0:
            kwargs = {"fontfile": str(hidden_font), "fontname": "hid"} if hidden_font else {"fontname": "helv"}
            new.insert_textbox(pymupdf.Rect(60, 60, 535, 500), hidden_text, fontsize=13, render_mode=3, **kwargs)
    out.save(path, deflate=True, garbage=3)


def _mojibake(text: str) -> str:
    """Mimics a legacy Indic font: glyphs mapped onto ASCII plus Latin-1 letters (about a third)."""
    out = []
    for ch in text:
        if ch.isspace():
            out.append(" ")
        elif ord(ch) % 10 < 3:
            out.append(chr(0xC0 + ord(ch) % 0x3F))
        else:
            out.append(chr(0x61 + ord(ch) % 26))
    return "".join(out)


def _private_use(text: str) -> str:
    font = pymupdf.Font(fontfile=str(SEGOE_ICONS))
    glyphs = [cp for cp in range(0xE700, 0xE800) if font.has_glyph(cp)]
    return "".join(" " if ch.isspace() else chr(glyphs[ord(ch) % len(glyphs)]) for ch in text)


def _patch_cached_values(path: Path, cached: dict[str, str]) -> None:
    """openpyxl writes formulas with an empty cached value; put real ones in, as Excel would."""
    with zipfile.ZipFile(path) as z:
        members = {n: z.read(n) for n in z.namelist()}
    sheet = members["xl/worksheets/sheet1.xml"].decode("utf-8")
    for cell, value in cached.items():
        sheet = re.sub(rf'(<c r="{cell}"[^>]*><f>[^<]*</f>)<v></v>', rf"\g<1><v>{value}</v>", sheet)
    members["xl/worksheets/sheet1.xml"] = sheet.encode("utf-8")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in members.items():
            z.writestr(name, data)


def _write_pptx(path: Path, picture_png: bytes) -> None:
    """A minimal two-slide deck written by hand (python-pptx is not a dependency)."""
    ns = (
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"'
    )
    rels_ns = "http://schemas.openxmlformats.org/package/2006/relationships"

    def shape(text, ph=""):
        return f'<p:sp><p:nvSpPr><p:cNvPr id="2" name="s"/><p:cNvSpPr/><p:nvPr>{ph}</p:nvPr></p:nvSpPr><p:txBody><a:p><a:r><a:t>{text}</a:t></a:r></a:p></p:txBody></p:sp>'

    def cell(t):
        return f"<a:tc><a:txBody><a:p><a:r><a:t>{t}</a:t></a:r></a:p></a:txBody></a:tc>"

    table = (
        "<p:graphicFrame><a:graphic><a:graphicData><a:tbl>"
        f"<a:tr>{cell('Item')}{cell('मद')}</a:tr><a:tr>{cell('Rice')}{cell('चावल')}</a:tr>"
        "</a:tbl></a:graphicData></a:graphic></p:graphicFrame>"
    )
    picture = '<p:pic><p:blipFill><a:blip r:embed="rId2"/></p:blipFill></p:pic>'
    title_placeholder = '<p:ph type="title"/>'
    slide1 = f"<p:sld {ns}><p:cSld><p:spTree>{shape(EN_TITLE, title_placeholder)}{shape(HI_P2)}{table}{picture}</p:spTree></p:cSld></p:sld>"
    slide2 = f"<p:sld {ns}><p:cSld><p:spTree>{shape('End of deck')}</p:spTree></p:cSld></p:sld>"
    parts = {
        "[Content_Types].xml": '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="xml" ContentType="application/xml"/></Types>',
        "_rels/.rels": f'<Relationships xmlns="{rels_ns}"><Relationship Id="rId1" Type="x" Target="ppt/presentation.xml"/></Relationships>',
        "ppt/presentation.xml": f'<p:presentation {ns}><p:sldIdLst><p:sldId id="256" r:id="rId1"/><p:sldId id="257" r:id="rId2"/></p:sldIdLst></p:presentation>',
        "ppt/_rels/presentation.xml.rels": f'<Relationships xmlns="{rels_ns}"><Relationship Id="rId1" Type="x" Target="slides/slide1.xml"/><Relationship Id="rId2" Type="x" Target="slides/slide2.xml"/></Relationships>',
        "ppt/slides/slide1.xml": slide1,
        "ppt/slides/_rels/slide1.xml.rels": f'<Relationships xmlns="{rels_ns}"><Relationship Id="rId2" Type="x" Target="../media/image1.png"/></Relationships>',
        "ppt/slides/slide2.xml": slide2,
        "ppt/media/image1.png": picture_png,
    }
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in parts.items():
            z.writestr(name, data)


def build(out: Path) -> dict:
    """Create all fixtures in ``out``. Returns the ground truth (also written to truth.json)."""
    out.mkdir(parents=True, exist_ok=True)
    truth: dict[str, dict] = {}

    # --- PDFs ---------------------------------------------------------------------------------
    en = english_doc()
    en.save(out / "en_born_digital.pdf")
    truth["en_born_digital.pdf"] = {"pages": [f"{EN_TITLE} {EN_P1}", EN_P2]}
    image_only(en, out / "en_image_only.pdf")
    truth["en_image_only.pdf"] = truth["en_born_digital.pdf"]
    image_only(en, out / "en_searchable_scan.pdf", hidden_text=f"{EN_TITLE} {EN_P1}")  # picture + exact hidden text
    truth["en_searchable_scan.pdf"] = {"pages": [f"{EN_TITLE} {EN_P1}", EN_P2]}

    hi_text_doc = pymupdf.open()  # exact Unicode text layer (unshaped glyph placement: the layer, not the look, is the point)
    page = hi_text_doc.new_page(width=595, height=842)
    page.insert_textbox(pymupdf.Rect(60, 60, 535, 400), f"{HI_TITLE}\n{HI_P1}\n{HI_P2}", fontsize=13, fontname="nir", fontfile=str(NIRMALA))
    hi_text_doc.save(out / "hi_born_digital.pdf")
    truth["hi_born_digital.pdf"] = {"pages": [HI_PAGE_TEXT]}

    hi = hindi_shaped_doc()
    hi.save(out / "hi_shaped_mupdf.pdf")  # looks right, but its text layer is damaged (see progress.txt)
    truth["hi_shaped_mupdf.pdf"] = {"pages": [HI_PAGE_TEXT]}
    image_only(hi, out / "hi_image_only.pdf")
    truth["hi_image_only.pdf"] = {"pages": [HI_PAGE_TEXT]}
    image_only(hi, out / "hi_broken_pua.pdf", hidden_text=_private_use(HI_PAGE_TEXT), hidden_font=SEGOE_ICONS)
    truth["hi_broken_pua.pdf"] = {"pages": [HI_PAGE_TEXT]}
    image_only(hi, out / "hi_broken_legacy.pdf", hidden_text=_mojibake(HI_PAGE_TEXT))
    truth["hi_broken_legacy.pdf"] = {"pages": [HI_PAGE_TEXT]}

    degraded = pymupdf.open()
    dp = degraded.new_page(width=595, height=842)
    dp.insert_image(dp.rect, stream=_degraded_jpeg(hi[0]))
    degraded.save(out / "hi_scan_degraded.pdf")
    truth["hi_scan_degraded.pdf"] = {"pages": [HI_PAGE_TEXT]}

    mixed = pymupdf.open()
    mixed.insert_pdf(en, from_page=0, to_page=0)
    mp = mixed.new_page(width=595, height=842)
    mp.insert_image(mp.rect, stream=_png(en[1], 200))
    mixed.save(out / "mixed_native_and_scan.pdf", deflate=True, garbage=3)
    truth["mixed_native_and_scan.pdf"] = {"pages": [f"{EN_TITLE} {EN_P1}", EN_P2]}

    blank = pymupdf.open()
    blank.new_page()
    blank.save(out / "blank.pdf")

    # --- Office and data files ----------------------------------------------------------------
    sentence_doc = _shaped(f"<p>{HI_SENTENCE}</p>", size_pt=16, width=420, height=60, box=(10, 5, 410, 55))
    sentence_png = _png(sentence_doc[0], 200)
    truth["docx_image"] = {"pages": [HI_SENTENCE]}

    import docx
    from docx.shared import Inches

    d = docx.Document()
    d.add_heading(EN_TITLE, level=1)
    d.add_paragraph(EN_P1)
    d.add_paragraph(HI_P2)
    table = d.add_table(rows=3, cols=3)
    table.style = "Table Grid"
    for r, row in enumerate([["Item", "मद", "Qty"], ["Rice", "चावल", "5"], ["Tea", "चाय", "12"]]):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    d.add_paragraph("Scanned sentence follows:")
    d.add_picture(io.BytesIO(sentence_png), width=Inches(5))
    d.add_paragraph("End of document.")
    d.save(out / "mixed.docx")

    import openpyxl
    from openpyxl.drawing.image import Image as XlImage

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    for row in (["Item", "Qty", "Price", "Total"], ["Rice", 5, 40, "=B2*C2"], ["Tea", 12, 10, "=B3*C3"], ["", "", "Sum", "=SUM(D2:D3)"]):
        ws.append(row)
    ws2 = wb.create_sheet("हिंदी")
    for row in (["नाम", "शहर"], ["राम", "दिल्ली"], ["सीता", "मुंबई"]):
        ws2.append(row)
    wb.save(out / "book.xlsx")
    _patch_cached_values(out / "book.xlsx", {"D2": "200", "D3": "120"})  # D4 stays uncached on purpose

    wb2 = openpyxl.Workbook()
    wb2.active.title = "Scan"
    wb2.active["A1"] = "Picture below"
    wb2.active.add_image(XlImage(io.BytesIO(sentence_png)), "A3")
    wb2.save(out / "book_image.xlsx")

    (out / "assets").mkdir(exist_ok=True)
    (out / "assets" / "sentence.png").write_bytes(sentence_png)
    _write_pptx(out / "deck.pptx", sentence_png)
    import base64

    data_uri = "data:image/png;base64," + base64.b64encode(sentence_png).decode("ascii")
    (out / "page.html").write_text(
        '<!doctype html><html><head><meta charset="utf-8"><title>Quarterly Report</title><script>var hidden = 1;</script></head><body>'
        f"<h1>Quarterly Report</h1><p>{EN_P2} <b>Bold</b> text.</p><p>{HI_P2}</p>"
        "<ul><li>one</li><li>दो</li></ul>"
        "<table><tr><th>Item</th><th>मद</th></tr><tr><td>Rice</td><td>चावल</td></tr></table>"
        '<img src="assets/sentence.png"><p>after the image</p>'
        f'<img src="https://example.com/remote.png"><img src="{data_uri}"></body></html>',
        encoding="utf-8",
    )

    (out / "table.csv").write_bytes(
        ("नाम,शहर,टिप्पणी\r\nराम,दिल्ली,\"पहली, दूसरी\"\r\nJohn,Mumbai,\"two\nlines\"\r\nx|y,\"a \"\"quoted\"\" word\",ok\r\n").encode("utf-8-sig")
    )
    (out / "nested.json").write_text(
        json.dumps(
            {"नाम": "राम", "शहर": {"नाम": "दिल्ली", "आबादी": 32000000}, "tags": ["भारत", "English", "日本語", "🙂"], "ok": True, "none": None},
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    (out / "note.xml").write_text(
        '<?xml version="1.0" encoding="UTF-8"?>\n<नोट id="7"><शीर्षक>नमस्ते दुनिया</शीर्षक><body lang="hi">यह एक परीक्षण है।</body></नोट>',
        encoding="utf-8",
    )
    (out / "हिंदी_पाठ.txt").write_text(f"{HI_TITLE}\n{HI_P1}\n", encoding="utf-8")
    (out / "plain.txt").write_bytes(b"first line\r\nsecond line\r\n")
    (out / "latin1_note.txt").write_bytes("caf\xe9 au lait".encode("cp1252"))

    # --- bad inputs ---------------------------------------------------------------------------
    (out / "corrupt.pdf").write_bytes(b"%PDF-1.4\n" + bytes(range(256)) * 20)
    (out / "empty.txt").write_bytes(b"")
    (out / "empty.pdf").write_bytes(b"")
    locked = pymupdf.open()
    locked.new_page().insert_text((60, 80), "secret")
    locked.save(out / "encrypted.pdf", encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="secret", owner_pw="owner")
    (out / "corrupt.docx").write_bytes(b"PK\x03\x04" + bytes(range(256)) * 4)
    (out / "broken.json").write_text('{"a": [1, 2,', encoding="utf-8")
    (out / "broken.xml").write_text("<a><b></a>", encoding="utf-8")

    (out / "truth.json").write_text(json.dumps(truth, ensure_ascii=False, indent=2), encoding="utf-8")
    return truth


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    target = Path(sys.argv[1])
    build(target)
    print(f"fixtures written to {target}")
